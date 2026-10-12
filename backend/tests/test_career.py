"""End-to-end tests for the career copilot API (/api/career) on an isolated SQLite DB.

LLM keys are cleared so every feature runs through its deterministic fallback."""
import os

os.environ["JOB_INGEST_ENABLED"] = "0"

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.career import autofill, enrichment, fraud, tools, vectors
from app.career.skills import extract_skills
from app.database import Base, get_db

RESUME = """Jane Doe
jane@example.com | +1 415 555 0182 | San Francisco, CA
Professional Summary
Machine learning engineer with 6 years of experience building NLP and data products.
Experience
Senior ML Engineer - Acme AI (Jan 2021 - Present)
• Responsible for building RAG pipelines with Python, PyTorch and LangChain serving 2M users
• Worked on fine-tuning transformers for document classification, improving accuracy by 18%
Data Scientist - DataCo (Mar 2018 - Dec 2020)
• Built feature engineering pipelines in Pandas and SQL for churn models
• Helped deploy models on GCP with Docker and Kubernetes
Skills
Python, PyTorch, NLP, SQL, Pandas, Docker, GCP, Machine Learning, Transformers, LangChain
Education
MS Computer Science - Stanford University (2016 - 2018)
"""
H = {"X-Client-Id": "pytest-client-0001"}


@pytest.fixture()
def client(monkeypatch):
    for key in ("GROQ_API_KEY", "OPENAI_API_KEY"):
        monkeypatch.setattr(f"app.config.settings.{key}", "", raising=False)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    for key in ("UPI_ID", "UPI_QR_IMAGE", "ADMIN_TOKEN", "ADMIN_EMAIL",   # default to dev billing
                "SMS_WEBHOOK_SECRET", "SMS_ALLOWED_SENDERS"):
        monkeypatch.delenv(key, raising=False)
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def _db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    import main
    vectors.job_index.invalidate()
    main.app.dependency_overrides[get_db] = _db
    with TestClient(main.app) as c:
        yield c
    main.app.dependency_overrides.pop(get_db, None)
    vectors.job_index.invalidate()


def _setup_profile(client, **prefs):
    r = client.put("/api/career/me", headers=H, json={"full_name": "Jane Doe", "email": "jane@example.com",
                                                      "resume_text": RESUME, "preferences": prefs})
    assert r.status_code == 200


def _first_job(client):
    jobs = client.get("/api/career/jobs?limit=5", headers=H).json()["jobs"]
    assert jobs
    return jobs[0]


# ------------------------------------------------------------------ units --

def test_skill_extraction_uses_word_boundaries():
    found = extract_skills("Senior engineer building a search engine with LangChain and Python")
    assert "langchain" in found and "python" in found
    assert "r" not in found and "gin" not in found and "chai" not in found


def test_h1b_signals():
    assert enrichment.h1b_status("Google LLC", "Build things.") == (True, "history")
    assert enrichment.h1b_status("Tiny Startup", "Visa sponsorship is available.") == (True, "stated")
    assert enrichment.h1b_status("Google", "We are unable to sponsor visas.") == (False, "denied")
    assert enrichment.h1b_status("Tiny Startup", "Great role.") == (False, "unknown")


def test_verification_flags_scams():
    ok, flags = enrichment.verify_listing("Data Entry", "ACME", "Earn $500 per day! Pay for training kit via gift card payment. " * 5,
                                          "https://bit.ly/x", None)
    assert not ok and any("pay" in f.lower() for f in flags)


def test_salary_parsing():
    assert enrichment.parse_salary("Base pay $150,000 - $190,000 per year") == (150000, 190000)
    assert enrichment.parse_salary("$120k-$160k") == (120000, 160000)
    assert enrichment.parse_salary("no salary listed") == (None, None)


def test_local_embedding_is_semantic():
    a = vectors.local_embedding("python machine learning engineer pytorch nlp")
    b = vectors.local_embedding("machine learning engineer using pytorch and python for nlp")
    c = vectors.local_embedding("registered nurse hospital patient care")
    import numpy as np
    assert vectors.cosine(np.array(a), np.array(b)) > vectors.cosine(np.array(a), np.array(c)) + 0.2


def test_rule_tailor_never_invents_skills():
    result = tools.tailor_resume(RESUME, "Kubernetes, Rust and Python engineer", "Engineer", "Co")
    assert result["engine"] == "rules"
    assert "rust" in result["missing_keywords"]
    assert all("Rust" not in b["optimized"] for b in result["bullets"])
    assert result["bullets"][0]["optimized"].startswith("Built RAG pipelines")


def test_fraud_detection_flags_fake_candidate():
    fake = "John Smith\njohn@mailinator.com\n111-111-1111\nCEO 2023 - 2024\nIntern 2022 - 2023\nIncreased revenue by 900%\n"
    result = fraud.check_candidate(fake, name="John Smith", use_llm=False)
    assert result["verdict"] == "high_risk"
    cats = {f["category"] for f in result["flags"]}
    assert {"identity", "claims"} <= cats


def test_autofill_mapping_rules():
    from app.career.models import SeekerProfile
    profile = SeekerProfile(client_id="x" * 10, full_name="Jane Doe", email="jane@example.com",
                            preferences='{"needs_sponsorship": true, "work_authorized": true}', resume_text=RESUME)
    fields = [
        {"id": "a", "tag": "input", "name": "firstName"},
        {"id": "b", "tag": "input", "label": "Last name"},
        {"id": "c", "tag": "select", "label": "Will you require visa sponsorship?", "options": ["Yes", "No"]},
        {"id": "d", "tag": "select", "label": "Veteran status", "options": ["I am a veteran", "I don't wish to answer"]},
        {"id": "e", "tag": "input", "type": "file", "label": "Resume"},
    ]
    fills = {f["id"]: f["value"] for f in autofill.map_fields(profile, fields)["fills"]}
    assert fills == {"a": "Jane", "b": "Doe", "c": "Yes", "d": "I don't wish to answer"}


# ------------------------------------------------------------------ API ----

def test_requires_client_id(client):
    assert client.get("/api/career/me").status_code == 400


def test_job_matching_and_filters(client):
    _setup_profile(client, roles=["ML Engineer"])
    data = client.get("/api/career/jobs?limit=10", headers=H).json()
    assert data["total"] > 20
    scores = [j["match"]["score"] for j in data["jobs"]]
    assert scores == sorted(scores, reverse=True)
    assert "ML" in data["jobs"][0]["title"] or "Machine Learning" in data["jobs"][0]["title"]
    h1b = client.get("/api/career/jobs?h1b_only=true", headers=H).json()["jobs"]
    assert h1b and all(j["h1b_sponsor"] for j in h1b)


def test_swipe_like_adds_to_tracker_and_hides_from_feed(client):
    _setup_profile(client)
    job = _first_job(client)
    r = client.post(f"/api/career/jobs/{job['id']}/decision", headers=H, json={"decision": "like"})
    assert r.json()["application"]["stage"] == "Wishlist"
    feed_ids = [j["id"] for j in client.get("/api/career/jobs?feed=true&limit=100", headers=H).json()["jobs"]]
    assert job["id"] not in feed_ids
    apps = client.get("/api/career/applications", headers=H).json()["applications"]
    assert len(apps) == 1
    patched = client.patch(f"/api/career/applications/{apps[0]['id']}", headers=H, json={"stage": "Applied"}).json()
    assert patched["stage"] == "Applied" and patched["applied_at"]


def test_free_plan_credits_and_upgrade(client):
    _setup_profile(client)
    job = _first_job(client)
    assert client.post("/api/career/resume/tailor", headers=H, json={"job_id": job["id"]}).status_code == 200
    assert client.post("/api/career/cover-letter", headers=H, json={"job_id": job["id"]}).status_code == 200
    assert client.post("/api/career/cover-letter", headers=H, json={"job_id": job["id"]}).status_code == 402
    status = client.post("/api/career/plan/upgrade", headers=H, json={}).json()
    assert status["plan"] == "turbo" and status["credits_limit"] is None
    assert client.post("/api/career/cover-letter", headers=H, json={"job_id": job["id"]}).status_code == 200


def test_pricing_periods_and_expiry(client):
    from datetime import datetime
    data = client.get("/api/career/plans").json()
    turbo = next(p for p in data["plans"] if p["id"] == "turbo")
    assert turbo["price_usd"] == 39.99
    periods = {p["id"]: p for p in data["billing_periods"]}
    assert set(periods) == {"monthly", "quarterly", "annual"}
    assert periods["annual"]["per_month_usd"] < periods["quarterly"]["per_month_usd"] < 39.99
    assert periods["annual"]["savings_pct"] > 0
    assert client.post("/api/career/plan/upgrade", headers=H, json={"period": "weekly"}).status_code == 422
    status = client.post("/api/career/plan/upgrade", headers=H, json={"period": "annual"}).json()
    days = (datetime.fromisoformat(status["expires_at"]) - datetime.utcnow()).days
    assert 363 <= days <= 365
    client.post("/api/career/plan/cancel", headers=H)
    status = client.post("/api/career/plan/upgrade", headers=H, json={"period": "quarterly"}).json()
    assert 88 <= (datetime.fromisoformat(status["expires_at"]) - datetime.utcnow()).days <= 90
    assert client.post("/api/career/plan/checkout", headers=H, json={"period": "annual"}).json()["mode"] == "dev"


def test_paypal_capture_validation():
    from app.career import plans
    cid = "pytest-client-0001"

    def cap(custom, value, status="COMPLETED"):
        return {"status": status, "purchase_units": [{"custom_id": custom,
                "payments": {"captures": [{"amount": {"currency_code": "USD", "value": value}}]}}]}

    assert plans.period_from_capture(cap(f"career|{cid}|annual", "359.99"), cid) == "annual"
    assert plans.period_from_capture(cap(f"career|{cid}", "39.99"), cid) == "monthly"       # legacy order
    assert plans.period_from_capture(cap(f"career|{cid}|annual", "39.99"), cid) is None     # underpaid
    assert plans.period_from_capture(cap(f"career|other|annual", "359.99"), cid) is None    # wrong owner
    assert plans.period_from_capture(cap(f"career|{cid}|annual", "359.99", "PENDING"), cid) is None
    assert plans.period_from_capture({}, cid) is None


def test_auto_apply_builds_packet(client):
    _setup_profile(client)
    job = _first_job(client)
    result = client.post("/api/career/auto-apply", headers=H, json={"job_id": job["id"]}).json()
    assert result["status"] == "ready"
    app_id = result["application"]["id"]
    packet = client.get(f"/api/career/applications/{app_id}/packet", headers=H).json()
    assert packet["tailored_resume"] and packet["cover_letter"]


def test_orion_rules_routing(client):
    _setup_profile(client)
    reply = client.post("/api/career/copilot/chat", headers=H, json={"message": "find ml engineer jobs"}).json()
    assert reply["tool_calls"][0]["tool"] == "search_jobs"
    reply = client.post("/api/career/copilot/chat", headers=H, json={"message": "how should I negotiate salary?"}).json()
    assert reply["tool_calls"][0]["tool"] == "salary_guidance"
    history = client.get("/api/career/copilot/history", headers=H).json()["messages"]
    assert len(history) == 4


def test_insider_connections_are_turbo_only(client):
    _setup_profile(client)
    csv = (b"Notes:\nexport\n\nFirst Name,Last Name,URL,Email Address,Company,Position,Connected On\n"
           b"Sam,Lee,https://linkedin.com/in/sam,,Google,Staff Engineer,01 Jan 2024\n")
    assert client.post("/api/career/connections/import", headers=H,
                       files={"file": ("Connections.csv", csv, "text/csv")}).json()["added"] == 1
    google = next(j for j in client.get("/api/career/jobs?query=google&limit=5", headers=H).json()["jobs"])
    locked = client.get(f"/api/career/jobs/{google['id']}/insiders", headers=H).json()
    assert locked["count"] == 1 and locked["locked"] and locked["insiders"] == []
    client.post("/api/career/plan/upgrade", headers=H, json={})
    unlocked = client.get(f"/api/career/jobs/{google['id']}/insiders", headers=H).json()
    assert unlocked["insiders"][0]["name"] == "Sam Lee"


def test_alert_limit_on_free_plan(client):
    _setup_profile(client)
    assert client.post("/api/career/alerts", headers=H, json={"query": "engineer", "min_match": 0}).status_code == 200
    assert client.post("/api/career/alerts", headers=H, json={"query": "data"}).status_code == 402
    assert client.get("/api/career/alerts/hits", headers=H).json()["hits"]


def test_recruiter_flow_ranks_and_blocks_fraud(client):
    _setup_profile(client)
    client.put("/api/career/me", headers=H, json={"open_to_recruiters": True})
    role = client.post("/api/career/recruiter/roles", headers=H, json={
        "title": "Senior ML Engineer", "company": "Foo",
        "description": "We need an ML engineer with Python, PyTorch, NLP, LangChain and RAG experience."}).json()
    sourced = client.post(f"/api/career/recruiter/roles/{role['id']}/source", headers=H).json()["added"]
    assert sourced and sourced[0]["name"] == "Jane Doe" and sourced[0]["fraud_verdict"] == "verified"
    fake = "John Smith\njohn@mailinator.com\n111-111-1111\nCEO 2023 - 2024\nIntern 2022 - 2023\nIncreased revenue by 900%\n"
    added = client.post(f"/api/career/recruiter/roles/{role['id']}/candidates", headers=H,
                        data={"resume_text": fake, "name": "John Smith"}).json()["added"][0]
    assert added["fraud_verdict"] == "high_risk"
    blocked = client.post(f"/api/career/recruiter/candidates/{added['id']}/contact", headers=H,
                          json={"subject": "Hi", "body": "Hello there, let's talk."})
    assert blocked.status_code == 409
    other = {"X-Client-Id": "someone-else-999"}
    assert client.get(f"/api/career/recruiter/roles/{role['id']}", headers=other).status_code == 404


# ------------------------------------------------------- manual UPI ----

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
ADMIN = {"X-Admin-Token": "test-admin-token-123"}


@pytest.fixture()
def upi(client, monkeypatch, tmp_path):
    from app.career import payments
    monkeypatch.setenv("UPI_ID", "profileoptimizer@ybl")
    monkeypatch.setenv("UPI_PAYEE_NAME", "Profile Optimizer")
    monkeypatch.setenv("ADMIN_TOKEN", ADMIN["X-Admin-Token"])
    monkeypatch.setenv("ADMIN_EMAIL", "admin@example.com")
    monkeypatch.setattr("app.config.settings.UPLOAD_DIR", str(tmp_path))
    sent = []
    monkeypatch.setattr(payments, "_notify", lambda to, subject, body: sent.append((to, subject)))
    client.sent = sent
    return client


def _submit(client, utr="123456789012", period="monthly", headers=H, image=PNG, name="proof.png"):
    return client.post("/api/career/payments", headers=headers,
                       data={"period": period, "utr": utr, "email": "jane@example.com"},
                       files={"screenshot": (name, image, "image/png")})


def test_upi_mode_details_and_create_request(upi, tmp_path):
    plans_data = upi.get("/api/career/plans").json()
    assert plans_data["billing_mode"] == "upi_manual"
    assert {p["id"]: p["price_inr"] for p in plans_data["billing_periods"]} ==         {"monthly": 499, "quarterly": 1299, "annual": 3999}
    # Instant upgrade is closed in this mode: access only comes from an admin approval.
    assert upi.post("/api/career/plan/upgrade", headers=H, json={}).status_code == 409

    info = upi.get("/api/career/payments/upi?period=quarterly", headers=H).json()
    assert info["amount_inr"] == 1299 and info["reference_code"].startswith("PO-")
    assert "pa=profileoptimizer@ybl" in info["upi_link"] and "am=1299.00" in info["upi_link"]
    assert info["qr_data_uri"].startswith("data:image/svg+xml") and info["qr_image_url"] is None

    r = _submit(upi, period="quarterly")
    assert r.status_code == 200, r.text
    req = r.json()["request"]
    assert req["status"] == "pending" and req["amount_inr"] == 1299
    assert req["reference_code"] == info["reference_code"]
    assert len(list((tmp_path / "payments").iterdir())) == 1
    assert upi.sent == [("admin@example.com", upi.sent[0][1])]
    assert upi.get("/api/career/payments/me", headers=H).json()["request"]["id"] == req["id"]
    me = upi.get("/api/career/me", headers=H).json()
    assert me["plan"]["plan"] == "free" and me["plan"]["payment"]["status"] == "pending"
    # One pending request at a time.
    assert _submit(upi, utr="999999999999").status_code == 409


def test_upi_request_validation(upi):
    assert _submit(upi, utr="12345").status_code == 422
    assert _submit(upi, image=b"GIF89a" + b"\x00" * 20, name="proof.png").status_code == 415
    assert _submit(upi, image=b"\xff\xd8\xff" + b"\x00" * (5 * 1024 * 1024)).status_code == 413


def test_upi_duplicate_utr_rejected(upi):
    assert _submit(upi).status_code == 200
    other = {"X-Client-Id": "pytest-client-0002"}
    r = _submit(upi, headers=other)
    assert r.status_code == 409 and "UTR" in r.json()["detail"]


def test_upi_admin_auth_required(upi):
    assert upi.get("/api/career/admin/payments").status_code == 401
    assert upi.get("/api/career/admin/payments", headers={"X-Admin-Token": "wrong"}).status_code == 401
    rid = _submit(upi).json()["request"]["id"]
    assert upi.post(f"/api/career/admin/payments/{rid}/approve", headers=H).status_code == 401
    assert upi.get(f"/api/career/admin/payments/{rid}/screenshot").status_code == 401
    shot = upi.get(f"/api/career/admin/payments/{rid}/screenshot", headers=ADMIN)
    assert shot.status_code == 200 and shot.content == PNG


def test_upi_approve_grants_turbo_once(upi):
    from datetime import datetime
    rid = _submit(upi, period="annual").json()["request"]["id"]
    pending = upi.get("/api/career/admin/payments?status=pending", headers=ADMIN).json()["requests"]
    assert [p["id"] for p in pending] == [rid]

    r = upi.post(f"/api/career/admin/payments/{rid}/approve", headers=ADMIN)
    assert r.status_code == 200 and r.json()["already_processed"] is False
    plan = upi.get("/api/career/me", headers=H).json()["plan"]
    assert plan["plan"] == "turbo" and plan["payment"]["status"] == "approved"
    expires = plan["expires_at"]
    assert 364 <= (datetime.fromisoformat(expires) - datetime.utcnow()).days <= 365
    assert ("jane@example.com", "Your ProfileOptimizer Turbo plan is active") in upi.sent

    # A second click is a no-op, not a second year.
    again = upi.post(f"/api/career/admin/payments/{rid}/approve", headers=ADMIN)
    assert again.status_code == 200 and again.json()["already_processed"] is True
    assert upi.get("/api/career/me", headers=H).json()["plan"]["expires_at"] == expires
    assert upi.post(f"/api/career/admin/payments/{rid}/reject", headers=ADMIN,
                    json={"reason": "too late"}).status_code == 409

    # A renewal while still active stacks on the current expiry.
    rid2 = _submit(upi, utr="210987654321", period="monthly").json()["request"]["id"]
    upi.post(f"/api/career/admin/payments/{rid2}/approve", headers=ADMIN)
    renewed = upi.get("/api/career/me", headers=H).json()["plan"]["expires_at"]
    assert (datetime.fromisoformat(renewed) - datetime.fromisoformat(expires)).days == 30


def test_upi_reject_keeps_free(upi):
    rid = _submit(upi).json()["request"]["id"]
    assert upi.post(f"/api/career/admin/payments/{rid}/reject", headers=ADMIN, json={"reason": ""}).status_code == 422
    r = upi.post(f"/api/career/admin/payments/{rid}/reject", headers=ADMIN,
                 json={"reason": "UTR not found in bank statement"})
    assert r.status_code == 200
    plan = upi.get("/api/career/me", headers=H).json()["plan"]
    assert plan["plan"] == "free"
    assert plan["payment"]["status"] == "rejected"
    assert plan["payment"]["reject_reason"] == "UTR not found in bank statement"
    assert upi.post(f"/api/career/admin/payments/{rid}/approve", headers=ADMIN).status_code == 409
    rejected = upi.get("/api/career/admin/payments?status=rejected", headers=ADMIN).json()["requests"]
    assert [x["id"] for x in rejected] == [rid]
    # After a rejection the user may submit again.
    assert _submit(upi, utr="555555555555").status_code == 200


def test_upi_rate_limit(upi):
    from app.career import payments
    for i in range(payments.DAILY_REQUEST_LIMIT):
        rid = _submit(upi, utr=f"10000000000{i}").json()["request"]["id"]
        upi.post(f"/api/career/admin/payments/{rid}/reject", headers=ADMIN, json={"reason": "test reject"})
    assert _submit(upi, utr="100000000009").status_code == 429


# --------------------------------------------------- background agent mode --

def _db_session():
    import main
    return next(main.app.dependency_overrides[get_db]())


def test_agent_mode_prepares_and_waits_for_review(client):
    _setup_profile(client)
    assert client.put("/api/career/agent", headers=H, json={"enabled": True}).status_code == 402   # Turbo only
    assert client.put("/api/career/agent", headers=H, json={"daily_cap": 500}).status_code == 422
    client.post("/api/career/plan/upgrade", headers=H, json={})
    overview = client.put("/api/career/agent", headers=H,
                          json={"enabled": True, "min_match": 0, "daily_cap": 3, "auto_submit": False}).json()
    assert overview["settings"]["enabled"] and "Phone" in overview["missing_answers"]
    run = client.post("/api/career/agent/run", headers=H).json()
    assert run["stats"]["queued"] == 3 and run["today"]["by_status"]["needs_review"] == 3
    assert client.post("/api/career/agent/run", headers=H).json()["stats"]["queued"] == 0   # daily cap
    tasks = client.get("/api/career/agent/tasks", headers=H).json()["tasks"]
    assert len(tasks) == 3 and all(t["status"] == "needs_review" for t in tasks)
    assert "Phone" in tasks[0]["reason"]
    detail = client.get(f"/api/career/agent/tasks/{tasks[0]['id']}", headers=H).json()
    assert [e["step"] for e in detail["log"]][:2] == ["queued", "verify"]
    # The browser channel can't close a task that is waiting for the user.
    assert client.post(f"/api/career/agent/tasks/{tasks[0]['id']}/result", headers=H,
                       json={"outcome": "submitted"}).status_code == 409
    done = client.post(f"/api/career/agent/tasks/{tasks[0]['id']}/resolve", headers=H, json={"outcome": "submitted"}).json()
    assert done["status"] == "submitted"
    app = next(a for a in client.get("/api/career/applications", headers=H).json()["applications"]
               if a["id"] == done["application_id"])
    assert app["stage"] == "Applied"
    other = {"X-Client-Id": "someone-else-999"}
    assert client.get(f"/api/career/agent/tasks/{tasks[1]['id']}", headers=other).status_code == 404


def test_agent_mode_skips_unverified_and_hands_off_supported_ats(client):
    from app.career.models import PlatformJob
    client.put("/api/career/me", headers=H, json={
        "full_name": "Jane Doe", "email": "jane@example.com", "phone": "+1 415 555 0182", "resume_text": RESUME,
        "preferences": {"work_authorized": True, "needs_sponsorship": False}})
    client.post("/api/career/plan/upgrade", headers=H, json={})
    jobs = client.get("/api/career/jobs?limit=3", headers=H).json()["jobs"]
    db = _db_session()
    gh, lever_bad, other = (db.get(PlatformJob, j["id"]) for j in jobs)
    gh.ats_platform = "greenhouse"
    lever_bad.ats_platform, lever_bad.verified = "lever", False
    other.ats_platform = "workday"
    db.commit()
    client.put("/api/career/agent", headers=H, json={"enabled": True, "auto_submit": True, "min_match": 0, "daily_cap": 2})
    client.post("/api/career/agent/run", headers=H)
    tasks = {t["job_id"]: t for t in client.get("/api/career/agent/tasks", headers=H).json()["tasks"]}
    assert lever_bad.id not in tasks                       # unverified postings are never picked
    assert tasks[gh.id]["status"] == "prepared" and tasks[gh.id]["submit_channel"] == "extension"
    assert tasks[other.id]["status"] == "needs_review"     # unsupported ATS: manual submit
    subs = client.get("/api/career/agent/submissions", headers=H).json()["submissions"]
    assert [s["id"] for s in subs] == [tasks[gh.id]["id"]] and subs[0]["packet"]["cover_letter"]
    failed = client.post(f"/api/career/agent/tasks/{subs[0]['id']}/result", headers=H,
                         json={"outcome": "failed", "detail": "Captcha"}).json()
    assert failed["status"] == "needs_review" and failed["submit_channel"] == ""
    db.close()


# ------------------------------------------------------ company insights --

def test_company_insights_endpoint_and_orion_tool(client):
    from app.career.models import PlatformJob
    _setup_profile(client)
    job = _first_job(client)
    db = _db_session()
    row = db.get(PlatformJob, job["id"])
    row.salary_min, row.salary_max = 150000, 200000
    db.commit()
    company = job["company"]
    data = client.get(f"/api/career/companies/insights?name={company}", headers=H).json()
    assert data["company"] == company and data["open_roles"]["count"] >= 1
    assert data["salary"]["overall"]["max"] >= 200000
    assert data["skills"] and data["velocity"]["last_30_days"] >= 1
    assert data["engine"] == "rules" and company in data["summary"]
    assert data["insiders"] == {"count": 0, "locked": False, "people": []}
    empty = client.get("/api/career/companies/insights?name=Nonexistent%20Corp%20XYZ", headers=H).json()
    assert empty["open_roles"]["count"] == 0 and "haven't indexed" in empty["summary"]
    assert client.get("/api/career/companies/insights?name=x", headers=H).status_code == 422
    reply = client.post("/api/career/copilot/chat", headers=H, json={"message": f"tell me about {company}"}).json()
    assert reply["tool_calls"][0]["tool"] == "company_insights"
    assert reply["tool_calls"][0]["args"]["company"] == company
    # No company named: falls through to the tracker instead of guessing.
    reply = client.post("/api/career/copilot/chat", headers=H, json={"message": "tell me about my applications"}).json()
    assert reply["tool_calls"][0]["tool"] == "tracker_summary"
    db.close()


def test_upi_static_qr_image(upi, monkeypatch, tmp_path):
    qr = tmp_path / "qr.png"
    qr.write_bytes(PNG)
    monkeypatch.setenv("UPI_QR_IMAGE", str(qr))
    info = upi.get("/api/career/payments/upi", headers=H).json()
    assert info["qr_image_url"] == "/api/career/payments/qr-image" and info["qr_data_uri"] is None
    assert upi.get(info["qr_image_url"]).content == PNG


# ------------------------------------------------- UPI SMS auto-verify ----

SMS_SECRET = {"X-Webhook-Secret": "sms-secret-123"}


@pytest.mark.parametrize("message, expected", [
    ("Money Received - INR 499.00 in your HDFC Bank A/c xx1234 on 10-10-26 by A/c linked to VPA "
     "abc@okaxis (UPI Ref No 412345678901)", (49900, "412345678901")),
    ("Dear UPI user A/C X1234 credited by 499.0 on date 10Oct26 trf from JANE Refno 412345678901. "
     "If not u? call 1800111109. -SBI", (49900, "412345678901")),
    ("ICICI Bank Account XX123 credited:Rs. 1,299.00 on 10-Oct-26. Info UPI-412345678901-JANE. "
     "Available Balance is Rs. 10,000.00", (129900, "412345678901")),
    ("INR 499.00 credited to A/c no. XX1234 on 10-10-26 by UPI/P2A/412345678901/JANE - Axis Bank", (49900, "412345678901")),
    ("Rs 499.00 debited from A/c XX1234 UPI Ref 412345678901", None),
    ("Rs 499 credited to your account", None),
])
def test_parse_credit_sms(message, expected):
    from app.career import payments
    assert payments.parse_credit_sms(message) == expected


@pytest.fixture()
def upi_auto(upi, monkeypatch):
    monkeypatch.setenv("SMS_WEBHOOK_SECRET", SMS_SECRET["X-Webhook-Secret"])
    monkeypatch.setenv("SMS_ALLOWED_SENDERS", "HDFCBK")
    return upi


def _sms(client, utr="123456789012", amount="499.00", sender="AD-HDFCBK", headers=SMS_SECRET):
    return client.post("/api/career/payments/sms-webhook", headers=headers, json={
        "from": sender, "message": f"Money Received - INR {amount} in your HDFC Bank A/c xx1234 (UPI Ref No {utr})"})


def test_sms_webhook_auth_and_sender(upi_auto, monkeypatch):
    assert _sms(upi_auto, headers={}).status_code == 401
    assert _sms(upi_auto, headers={"X-Webhook-Secret": "wrong"}).status_code == 401
    assert _sms(upi_auto, sender="+919999999999").json()["status"] == "ignored"
    monkeypatch.delenv("SMS_ALLOWED_SENDERS")
    assert _sms(upi_auto).status_code == 404   # off without a sender allowlist


def test_sms_after_submit_auto_approves(upi_auto):
    assert upi_auto.get("/api/career/payments/upi", headers=H).json()["auto_verify"] is True
    r = upi_auto.post("/api/career/payments", headers=H,
                      data={"period": "monthly", "utr": "123456789012", "email": "jane@example.com"})
    assert r.status_code == 200 and r.json()["request"]["status"] == "pending"   # no screenshot needed

    assert _sms(upi_auto).json()["status"] == "matched"
    plan = upi_auto.get("/api/career/me", headers=H).json()["plan"]
    assert plan["plan"] == "turbo" and plan["payment"]["status"] == "approved"
    expires = plan["expires_at"]
    assert _sms(upi_auto).json()["status"] == "stored"   # a repeated SMS doesn't extend twice
    assert upi_auto.get("/api/career/me", headers=H).json()["plan"]["expires_at"] == expires


def test_sms_before_submit_auto_approves(upi_auto):
    assert _sms(upi_auto, utr="210987654321").json()["status"] == "stored"
    r = _submit(upi_auto, utr="210987654321")
    assert r.json()["request"]["status"] == "approved"


def test_sms_underpaid_stays_pending(upi_auto):
    _submit(upi_auto)
    assert _sms(upi_auto, amount="49.00").json()["status"] == "stored"
    assert upi_auto.get("/api/career/payments/me", headers=H).json()["request"]["status"] == "pending"


def test_screenshot_still_required_without_auto_verify(upi):
    r = upi.post("/api/career/payments", headers=H,
                 data={"period": "monthly", "utr": "123456789012", "email": "jane@example.com"})
    assert r.status_code == 422


# -------------------------------------------------------- resume builder --

def test_resume_builder_parse_save_and_match_source(client):
    empty = client.get("/api/career/resume/structured", headers=H).json()
    assert empty["source"] == "empty" and empty["resume"]["experience"] == []
    assert client.post("/api/career/resume/structured/parse", headers=H, data={"resume_text": "too short"}).status_code == 400
    bad = client.post("/api/career/resume/structured/parse", headers=H,
                      files={"file": ("cv.exe", b"MZ" * 100, "application/octet-stream")})
    assert bad.status_code == 415
    parsed = client.post("/api/career/resume/structured/parse", headers=H,
                         files={"file": ("cv.txt", RESUME.encode(), "text/plain")}).json()
    doc = parsed["resume"]
    assert parsed["source"] == "parsed" and doc["fullName"] == "Jane Doe" and doc["email"] == "jane@example.com"
    assert [e["company"] for e in doc["experience"]] == ["Acme AI", "DataCo"]
    assert doc["experience"][0]["startDate"] == "Jan 2021" and doc["experience"][0]["endDate"] == "Present"
    assert "PyTorch" in doc["skills"][0]["items"]
    assert doc["education"][0]["school"] == "Stanford University"
    # Parsing alone doesn't save anything.
    assert client.get("/api/career/me", headers=H).json()["has_resume"] is False
    doc["title"] = "Senior ML Engineer"
    doc["experience"][0]["bullets"][0]["optimized"] = "Built RAG pipelines in Python serving 2M users"
    doc["experience"][0]["bullets"][0]["isOptimized"] = True
    saved = client.put("/api/career/resume/structured", headers=H, json=doc).json()
    assert saved["source"] == "saved" and saved["profile"]["has_resume"] and saved["profile"]["headline"] == "Senior ML Engineer"
    # The structured resume is now what matching runs on.
    jobs = client.get("/api/career/jobs?limit=3", headers=H).json()["jobs"]
    assert jobs[0]["match"]["score"] > 0
    txt = client.get("/api/career/resume/structured/export?format=txt", headers=H)
    assert "Built RAG pipelines in Python serving 2M users" in txt.text and "Responsible for" not in txt.text
    other = {"X-Client-Id": "someone-else-999"}
    assert client.get("/api/career/resume/structured", headers=other).json()["source"] == "empty"
    too_long = {**doc, "summary": "x" * 5000}
    assert client.put("/api/career/resume/structured", headers=H, json=too_long).status_code == 422


def test_resume_builder_exports_ats_safe_files(client):
    import io
    from docx import Document
    from pypdf import PdfReader
    assert client.get("/api/career/resume/structured/export", headers=H).status_code == 400
    client.post("/api/career/resume/structured/parse", headers=H, data={"resume_text": RESUME, "save": "true"})
    pdf = client.get("/api/career/resume/structured/export?format=pdf", headers=H)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    text = PdfReader(io.BytesIO(pdf.content)).pages[0].extract_text()
    assert "Jane Doe" in text and "EXPERIENCE" in text and "Acme AI" in text
    docx = client.get("/api/career/resume/structured/export?format=docx", headers=H)
    paragraphs = [p.text for p in Document(io.BytesIO(docx.content)).paragraphs]
    assert "Jane Doe" in paragraphs and "Experience" in paragraphs and not Document(io.BytesIO(docx.content)).tables
    assert client.get("/api/career/resume/structured/export?format=html", headers=H).status_code == 422


# ------------------------------------------------- job discovery / boards --

LONG_DESC = ("<p>We are hiring a backend engineer to build Python and Kubernetes services for our payments platform. "
             "You will design APIs, own reliability, mentor teammates and work closely with product and data teams "
             "across the company to ship features that customers love. Competitive salary and benefits.</p>")


def _mock_client(routes):
    import httpx

    def handler(request):
        for (method, fragment), payload in routes.items():
            if request.method == method and fragment in str(request.url):
                return httpx.Response(200, json=payload)
        return httpx.Response(404, json={})
    return httpx.Client(transport=httpx.MockTransport(handler))


def test_workday_and_smartrecruiters_connectors_ingest_and_dedupe(client):
    from app.career.ingestion import SmartRecruitersConnector, WorkdayConnector, board_lists, run_ingestion
    from app.career.models import CompanyBoard, PlatformJob
    wd_board = "acme.wd5.myworkdayjobs.com/acme/External"
    wd = WorkdayConnector([wd_board, "not-a-workday-board"], client=_mock_client({
        ("POST", "/wday/cxs/acme/External/jobs"): {"total": 1, "jobPostings": [
            {"title": "Backend Engineer", "externalPath": "/job/Austin-TX/Backend-Engineer_R1", "locationsText": "Austin, TX"}]},
        ("GET", "/wday/cxs/acme/External/job/Austin-TX/Backend-Engineer_R1"): {"jobPostingInfo": {
            "title": "Backend Engineer", "jobDescription": LONG_DESC, "location": "Austin, TX", "jobReqId": "R1",
            "startDate": "2026-10-01", "timeType": "Full time",
            "externalUrl": "https://acme.wd5.myworkdayjobs.com/External/job/Austin-TX/Backend-Engineer_R1"}},
    }))
    sr = SmartRecruitersConnector(["AcmeCo"], client=_mock_client({
        ("GET", "/companies/AcmeCo/postings/77"): {"postingUrl": "https://jobs.smartrecruiters.com/AcmeCo/77",
                                                    "jobAd": {"sections": {"jobDescription": {"text": LONG_DESC}}}},
        ("GET", "/companies/AcmeCo/postings"): {"totalFound": 2, "content": [
            {"id": "77", "name": "Data Analyst", "company": {"name": "AcmeCo"}, "releasedDate": "2026-10-02T00:00:00Z",
             "location": {"city": "Denver", "region": "CO", "country": "us", "remote": False},
             "typeOfEmployment": {"label": "Full-time"}},
            # Same role as the Workday one (same title/company/city): must dedupe, not duplicate.
            {"id": "78", "name": "Backend Engineer", "company": {"name": "Acme"}, "releasedDate": "2026-10-02T00:00:00Z",
             "location": {"city": "Austin", "region": "TX", "country": "us"}}]},
    }))
    db = _db_session()
    db.add(CompanyBoard(ats="workday", slug=wd_board, source="admin"))
    db.commit()
    stats = run_ingestion(db, [wd, sr])
    assert stats["new"] == 2 and stats["updated"] == 1
    jobs = {j.title: j for j in db.query(PlatformJob).filter(PlatformJob.source.in_(["workday", "smartrecruiters"])).all()}
    assert jobs["Backend Engineer"].ats_platform == "Workday" and jobs["Backend Engineer"].verified
    assert jobs["Data Analyst"].location == "Denver, CO, US" and jobs["Data Analyst"].url.endswith("/77")
    assert wd.board_stats["not-a-workday-board"]["error"]
    row = db.query(CompanyBoard).filter(CompanyBoard.slug == wd_board).one()
    assert row.last_status == "ok" and row.last_count == 1
    assert wd_board in board_lists(db)["workday"] and "stripe" in board_lists(db)["greenhouse"]
    db.close()


def test_admin_boards_api(client, monkeypatch):
    monkeypatch.setenv("ADMIN_TOKEN", "board-admin-token")
    A = {"X-Admin-Token": "board-admin-token"}
    assert client.get("/api/career/admin/boards").status_code == 401
    assert client.get("/api/career/admin/boards", headers={"X-Admin-Token": "wrong"}).status_code == 401
    boards = client.get("/api/career/admin/boards", headers=A).json()["boards"]
    stripe = next(b for b in boards if b["ats"] == "greenhouse" and b["slug"] == "stripe")
    assert stripe["source"] in ("default", "env")
    added = client.post("/api/career/admin/boards", headers=A,
                        json={"ats": "greenhouse", "slug": "https://boards.greenhouse.io/acmeco"}).json()["board"]
    assert added["slug"] == "acmeco" and added["source"] == "admin"
    assert client.post("/api/career/admin/boards", headers=A, json={"ats": "greenhouse", "slug": "AcmeCo"}).status_code == 409
    assert client.post("/api/career/admin/boards", headers=A, json={"ats": "lever", "slug": "bad slug!"}).status_code == 422
    assert client.post("/api/career/admin/boards", headers=A, json={"ats": "taleo", "slug": "x1"}).status_code == 422
    wd = client.post("/api/career/admin/boards", headers=A, json={
        "ats": "workday", "slug": "https://acme.wd5.myworkdayjobs.com/en-US/acme/External/"}).json()["board"]
    assert wd["slug"] == "acme.wd5.myworkdayjobs.com/acme/External"
    assert client.post("/api/career/admin/boards", headers=A,
                       json={"ats": "workday", "slug": "example.com/a/b"}).status_code == 422
    sr = client.post("/api/career/admin/boards", headers=A,
                     json={"ats": "smartrecruiters", "slug": "https://careers.smartrecruiters.com/AcmeCo"}).json()["board"]
    assert sr["slug"] == "AcmeCo"
    # Default boards are deactivated (they'd be re-seeded), admin-added ones are deleted.
    gone = client.delete(f"/api/career/admin/boards/{stripe['id']}", headers=A).json()
    assert gone["deactivated"] and not gone["board"]["active"]
    assert client.delete(f"/api/career/admin/boards/{added['id']}", headers=A).json()["deleted"]
    assert client.patch(f"/api/career/admin/boards/{wd['id']}", headers=A, json={"active": False}).json()["board"]["active"] is False
    from app.career.ingestion import board_lists
    db = _db_session()
    lists = board_lists(db)
    assert "stripe" not in lists["greenhouse"] and "acmeco" not in lists["greenhouse"]
    assert lists["workday"] == [] and lists["smartrecruiters"] == ["AcmeCo"]
    db.close()
    assert client.delete("/api/career/admin/boards/nope", headers=A).status_code == 404


# ------------------------------------- job APIs + browser-captured postings --

def test_adzuna_jooble_himalayas_connectors(client):
    import httpx
    from app.career.ingestion import AdzunaConnector, HimalayasConnector, JoobleConnector, run_ingestion
    from app.career.models import PlatformJob
    adzuna = AdzunaConnector("id", "key", countries=["in"], queries=["data analyst"], client=_mock_client({
        ("GET", "/jobs/in/search/1"): {"results": [{
            "id": 501, "title": "Data Analyst", "company": {"display_name": "Infosys"},
            "location": {"display_name": "Bengaluru, Karnataka"}, "description": LONG_DESC,
            "redirect_url": "https://www.adzuna.in/land/ad/501", "created": "2026-10-04T08:00:00Z",
            "salary_min": 900000, "salary_max": 1400000, "contract_time": "full_time"}]},
    }))

    def jooble_handler(request):
        assert request.method == "POST" and request.url.path == "/api/jkey"
        return httpx.Response(200, json={"jobs": [{"id": 9, "title": "Backend Developer", "company": "Zoho",
                                                    "location": "Chennai", "snippet": LONG_DESC,
                                                    "link": "https://jooble.org/desc/9", "updated": "2026-10-05T00:00:00"}]})
    jooble = JoobleConnector("jkey", queries=["backend"], locations=["India"],
                             client=httpx.Client(transport=httpx.MockTransport(jooble_handler)))
    himalayas = HimalayasConnector(max_jobs=20, client=_mock_client({
        ("GET", "himalayas.app/jobs/api"): {"jobs": [{"guid": "h1", "title": "Frontend Engineer", "companyName": "Acme Remote",
                                                      "locationRestrictions": ["India"], "description": LONG_DESC,
                                                      "applicationLink": "https://himalayas.app/jobs/h1", "pubDate": 1791000000}]},
    }))
    db = _db_session()
    stats = run_ingestion(db, [adzuna, jooble, himalayas])
    assert stats["new"] == 3 and stats["by_source"] == {"adzuna": 1, "jooble": 1, "himalayas": 1}
    da = db.query(PlatformJob).filter(PlatformJob.source == "adzuna").one()
    assert da.company == "Infosys" and da.salary_min == 900000 and da.url.startswith("https://www.adzuna.in/")
    assert db.query(PlatformJob).filter(PlatformJob.source == "himalayas").one().remote
    db.close()


def test_import_job_from_browser_is_private_and_tracked(client):
    _setup_profile(client)
    body = {"url": "https://www.linkedin.com/jobs/view/123", "title": "Senior ML Engineer", "company": "Flipkart",
            "location": "Bengaluru", "description": "We need Python, PyTorch and NLP experience to build ranking models. " * 3,
            "site": "linkedin.com", "salary_text": "$150,000 - $190,000"}
    r = client.post("/api/career/jobs/import", headers=H, json=body)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["existing"] is False and data["application"]["stage"] == "Wishlist"
    assert data["job"]["match"]["score"] > 0 and data["job"]["salary_min"] == 150000
    job_id = data["job"]["id"]
    # Usable by the importer (detail + AI tools) but not in anyone's shared feed or search.
    assert client.get(f"/api/career/jobs/{job_id}", headers=H).json()["title"] == "Senior ML Engineer"
    other = {"X-Client-Id": "someone-else-999"}
    ids = [j["id"] for j in client.get("/api/career/jobs?query=flipkart&limit=50", headers=other).json()["jobs"]]
    assert job_id not in ids
    # A second user importing the same posting doesn't make it public, and gets the same job.
    again = client.post("/api/career/jobs/import", headers=other, json=body).json()
    assert again["existing"] is True and again["job"]["id"] == job_id
    assert job_id not in [j["id"] for j in client.get("/api/career/jobs?query=flipkart&limit=50", headers=H).json()["jobs"]]
    # Importing a job we already crawl links to it and keeps it public.
    crawled = _first_job(client)
    linked = client.post("/api/career/jobs/import", headers=H, json={
        **body, "title": crawled["title"], "company": crawled["company"], "location": crawled["location"]}).json()
    assert linked["existing"] is True and linked["job"]["id"] == crawled["id"]
    visible = client.get("/api/career/jobs", headers=other, params={"query": crawled["company"], "limit": 50}).json()["jobs"]
    assert crawled["id"] in [j["id"] for j in visible]
    assert client.post("/api/career/jobs/import", headers=H, json={**body, "url": "javascript:alert(1)"}).status_code == 422


# ------------------------------------------------- improve after ATS check --

def test_resume_recheck_matches_upload_scorer_and_job(client):
    from app.ats_scorer import score_text
    assert client.post("/api/career/resume/recheck", headers=H, json={}).status_code == 400   # no resume yet
    _setup_profile(client)
    saved = client.post("/api/career/resume/recheck", headers=H, json={}).json()
    assert saved["ats_score"] == min(max(round(score_text(RESUME.strip())["unified_score"]), 0), 100)
    assert saved["job_match"] is None and isinstance(saved["issues"], list)
    edited = RESUME + "\n• Led a team of 5 engineers and increased revenue by 25% in 6 months\n"
    jd = "Senior ML Engineer: Python, PyTorch, NLP, Kubernetes, Spark and MLOps to ship LLM features."
    r = client.post("/api/career/resume/recheck", headers=H,
                    json={"resume_text": edited, "job_description": jd, "job_title": "ML Engineer"}).json()
    assert {"python", "pytorch", "nlp"} <= set(r["job_match"]["matched"]) and "spark" in r["job_match"]["missing"]
    assert all({"id", "label", "detail"} <= set(i) for i in r["issues"])
    assert client.post("/api/career/resume/recheck", headers=H, json={"resume_text": "x" * 60000}).status_code == 422


def test_parser_keeps_experience_without_bullet_characters():
    from app.career import resume_builder
    text = ("Sam Patel\nsam@example.com\nSummary\nI work on data stuff.\nExperience\nAnalyst at Acme 2021 - Present\n"
            "Responsible for reports in Excel and SQL\nHelped the team with dashboards\n"
            "Intern - Globex (Jun 2020 - Aug 2020)\nCleaned survey data in Python\nSkills\nExcel, SQL\n")
    doc = resume_builder.parse_text(text)
    roles = [(e.role, e.company, e.startDate, e.endDate, [b.original for b in e.bullets]) for e in doc.experience]
    assert roles == [
        ("Analyst", "Acme", "2021", "Present", ["Responsible for reports in Excel and SQL", "Helped the team with dashboards"]),
        ("Intern", "Globex", "Jun 2020", "Aug 2020", ["Cleaned survey data in Python"]),
    ]
    out = resume_builder.to_text(doc)
    assert "Helped the team with dashboards" in out and "Cleaned survey data in Python" in out
